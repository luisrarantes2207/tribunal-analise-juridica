# -*- coding: utf-8 -*-
"""Peca processual editavel (aba "Peca Processual" da v2).

- Reconstroi a peca do PDF (paragrafos, recuo, citacao em bloco, centralizado,
  negrito/italico/vermelho) ou le um DOCX enviado pelo usuario.
- Formato interno: HTML simples (<p>, <b>, <i>, <u>, <span style="color">),
  o mesmo que o editor da tela usa.
- Exporta DOCX sobre o papel timbrado (modelos/timbrado.docx: cabecalho, rodape
  e marca d'agua do escritorio) e PDF convertendo esse DOCX pelo Word.
- Assistente de IA que edita por operacoes sobre paragrafos numerados e
  avaliacao rapida da peca pelo mesmo roteiro de criterios do motor.
"""
import html
import json
import re
import subprocess
import tempfile
from collections import Counter
from html.parser import HTMLParser
from pathlib import Path

import extracao
import ia
import motor

BASE = Path(__file__).resolve().parent
TIMBRADO = BASE / "modelos" / "timbrado.docx"
VERMELHO = "#c00000"


# ======================================================== PDF -> blocos
def _negrito(fonte: str) -> bool:
    return bool(re.search(r"bold|black|heavy|semibold|demi", fonte or "", re.I))


def _italico(fonte: str) -> bool:
    return bool(re.search(r"italic|oblique", fonte or "", re.I))


def _cor(c) -> str | None:
    try:
        if isinstance(c, (list, tuple)) and len(c) == 3 and c[0] > 0.55 and c[1] < 0.35 and c[2] < 0.35:
            return VERMELHO
    except TypeError:
        pass
    return None


def _linhas_da_pagina(pg) -> list[dict]:
    palavras = [w for w in pg.extract_words(extra_attrs=["fontname", "size", "non_stroking_color"],
                                            keep_blank_chars=False, use_text_flow=False)
                if w.get("upright", True)]
    palavras.sort(key=lambda w: (round(w["top"] / 3), w["x0"]))
    linhas, atual = [], []
    for w in palavras:
        if atual and abs(w["top"] - atual[-1]["top"]) > 3:
            linhas.append(atual)
            atual = []
        atual.append(w)
    if atual:
        linhas.append(atual)
    out = []
    for ws in linhas:
        ws.sort(key=lambda w: w["x0"])
        runs = []
        for w in ws:
            fmt = (_negrito(w.get("fontname")), _italico(w.get("fontname")), _cor(w.get("non_stroking_color")))
            if runs and runs[-1]["fmt"] == fmt:
                runs[-1]["t"] += " " + w["text"]
            else:
                if runs:
                    runs[-1]["t"] += " "
                runs.append({"t": w["text"], "fmt": fmt})
        texto = " ".join(w["text"] for w in ws)
        out.append({"x0": ws[0]["x0"], "x1": ws[-1]["x1"], "top": min(w["top"] for w in ws),
                    "bottom": max(w["bottom"] for w in ws), "texto": texto, "runs": runs,
                    "size": Counter(round(w.get("size") or 12) for w in ws).most_common(1)[0][0]})
    return out


def _chave(t: str) -> str:
    return re.sub(r"\d+", "#", re.sub(r"\s+", " ", t.strip().lower()))


def pdf_para_blocos(caminho: str) -> list[dict]:
    import pdfplumber

    with pdfplumber.open(caminho) as pdf:
        paginas = [(_linhas_da_pagina(pg), float(pg.width)) for pg in pdf.pages]
    # cabecalho/rodape repetido (timbrado, numero de pagina, carimbo) sai da peca
    cont = Counter(_chave(l["texto"]) for ls, _ in paginas for l in {_chave(x["texto"]): x for x in ls}.values())
    repetidas = {k for k, v in cont.items() if len(paginas) >= 3 and v >= max(3, len(paginas) // 2)}
    linhas = []
    for ls, largura in paginas:
        for l in ls:
            if _chave(l["texto"]) in repetidas or re.fullmatch(r"(p[aá]g(ina)?\.?\s*)?\d+(\s*(de|/)\s*\d+)?", l["texto"].strip(), re.I):
                continue
            if extracao._eh_carimbo(l["texto"]):
                continue
            l["largura"] = largura
            linhas.append(l)
    if not linhas:
        return []
    esq = Counter(round(l["x0"]) for l in linhas).most_common(1)[0][0]
    dir_ = sorted(l["x1"] for l in linhas)[int(len(linhas) * 0.9)]
    alturas = sorted(b["top"] - a["bottom"] for a, b in zip(linhas, linhas[1:]) if 0 < b["top"] - a["bottom"] < 40)
    vao = alturas[len(alturas) // 2] if alturas else 6

    blocos, atual = [], None

    def fechar():
        if atual and atual["linhas"]:
            blocos.append(atual)

    for i, l in enumerate(linhas):
        ant = linhas[i - 1] if i else None
        novo = ant is None
        if ant is not None:
            mesma_pag = l["top"] > ant["top"]
            curta = ant["x1"] < dir_ - 30          # linha justificada so fica curta no fim do paragrafo
            gap = mesma_pag and (l["top"] - ant["bottom"]) > vao * 1.7 + 2
            recua = l["x0"] > esq + 12 and ant["x0"] <= esq + 6
            # bloco que comecou recuado: 2a linha na margem = recuo de 1a linha; recuada = citacao
            if len(atual["linhas"]) == 1 and atual["linhas"][0]["x0"] > esq + 12:
                atual["modo"] = "recuo" if l["x0"] <= esq + 6 else "citacao"
            sai_citacao = atual.get("modo") == "citacao" and l["x0"] <= esq + 6
            novo = gap or recua or sai_citacao or curta
        if novo:
            fechar()
            atual = {"linhas": []}
        atual["linhas"].append(l)
    fechar()
    return [_bloco(b["linhas"], esq, dir_) for b in blocos]


def _todo_negrito(l: dict) -> bool:
    return all(r["fmt"][0] for r in l["runs"] if r["t"].strip())


def _bloco(ls: list[dict], esq: float, dir_: float) -> dict:
    runs = []
    for k, l in enumerate(ls):
        for r in l["runs"]:
            t = r["t"]
            if runs and runs[-1]["fmt"] == r["fmt"]:
                runs[-1]["t"] += t
            else:
                runs.append({"t": t, "fmt": r["fmt"]})
        if k < len(ls) - 1:
            # hifenizacao: "123.456.789-" + "90" ou "manifes-" + "tacao"
            if runs[-1]["t"].endswith("-") and not runs[-1]["t"].endswith(" -"):
                continue
            runs[-1]["t"] += " "
    primeira = ls[0]
    largura = primeira["largura"]
    todas_recuadas = all(l["x0"] > esq + 40 for l in ls)
    if len(ls) <= 3 and all(abs((l["x0"] + l["x1"]) / 2 - largura / 2) < 18 for l in ls) and primeira["x0"] > esq + 25:
        classe, alinha = "", "center"
    elif all(abs(l["x1"] - dir_) < 8 for l in ls) and primeira["x0"] > esq + 120 and len(ls) <= 2:
        classe, alinha = "", "right"
    elif todas_recuadas and (len(ls) > 1 or primeira["x0"] > esq + 90):
        classe, alinha = "citacao", "justify"
    elif primeira["x0"] > esq + 12 and (len(ls) == 1 or ls[1]["x0"] <= esq + 6):
        classe, alinha = "recuo", "justify"
    else:
        classe, alinha = "", "justify" if len(ls) > 1 else "left"
    return {"classe": classe, "alinha": alinha, "runs": [r for r in runs if r["t"]]}


# ======================================================== DOCX -> blocos
def docx_para_blocos(caminho) -> list[dict]:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH as AL

    d = docx.Document(caminho)
    mapa = {AL.CENTER: "center", AL.RIGHT: "right", AL.JUSTIFY: "justify", AL.LEFT: "left"}
    out = []
    for p in d.paragraphs:
        runs = []
        for r in p.runs:
            if not r.text:
                continue
            cor = None
            try:
                if r.font.color is not None and r.font.color.rgb is not None:
                    rgb = str(r.font.color.rgb)
                    cor = "#" + rgb.lower() if rgb.upper() not in ("000000", "AUTO") else None
            except (ValueError, AttributeError):
                cor = None
            b = bool(r.bold) or bool(p.style is not None and p.style.font and p.style.font.bold)
            runs.append({"t": r.text, "fmt": (b, bool(r.italic), cor), "u": bool(r.underline)})
        pf = p.paragraph_format
        classe = ""
        if pf.left_indent and pf.left_indent.cm >= 2.5:
            classe = "citacao"
        elif pf.first_line_indent and pf.first_line_indent.cm >= 0.8:
            classe = "recuo"
        alinha = mapa.get(p.alignment if p.alignment is not None else (p.style.paragraph_format.alignment if p.style else None), "justify")
        if not runs:
            out.append({"classe": "", "alinha": "left", "runs": [], "vazio": True})
            continue
        out.append({"classe": classe, "alinha": alinha, "runs": runs})
    # colapsa linhas vazias repetidas
    limpo = []
    for b in out:
        if b.get("vazio") and limpo and limpo[-1].get("vazio"):
            continue
        limpo.append(b)
    return [b for b in limpo if not b.get("vazio")] if len(limpo) > 400 else limpo


# ======================================================== blocos <-> HTML
def blocos_para_html(blocos: list[dict]) -> str:
    partes = []
    for b in blocos:
        inner = ""
        for r in b["runs"]:
            t = html.escape(r["t"])
            neg, ita, cor = r["fmt"][:3]
            if cor:
                t = '<span style="color:%s">%s</span>' % (cor, t)
            if r.get("u"):
                t = "<u>%s</u>" % t
            if ita:
                t = "<i>%s</i>" % t
            if neg:
                t = "<b>%s</b>" % t
            inner += t
        attrs = ""
        if b.get("classe"):
            attrs += ' class="%s"' % b["classe"]
        if b.get("alinha") and b["alinha"] != "left":
            attrs += ' style="text-align:%s"' % b["alinha"]
        partes.append("<p%s>%s</p>" % (attrs, inner or "<br>"))
    return "\n".join(partes)


class _Leitor(HTMLParser):
    """HTML do editor -> blocos (paragrafos com trechos formatados)."""
    BLOCO = {"p", "div", "h1", "h2", "h3", "h4", "li", "blockquote"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.blocos, self.pilha, self.atual = [], [], None

    def _novo(self, tag, attrs):
        a = dict(attrs)
        estilo = a.get("style") or ""
        m = re.search(r"text-align\s*:\s*(\w+)", estilo)
        classe = a.get("class") or ""
        cls = "citacao" if "citacao" in classe or tag == "blockquote" else "recuo" if "recuo" in classe else ""
        if tag in ("h1", "h2", "h3", "h4"):
            cls = "titulo"
        self.atual = {"classe": cls, "alinha": m.group(1) if m else "", "runs": [], "tag": tag,
                      "lista": tag == "li"}

    def handle_starttag(self, tag, attrs):
        if tag in self.BLOCO:
            self._fechar()
            self._novo(tag, attrs)
            return
        if tag == "hr":
            self._fechar()
            self.blocos.append({"quebra": True})
            return
        if tag == "br":
            if self.atual is not None:
                self.atual["runs"].append({"t": "\n", "fmt": self._fmt()})
            return
        a = dict(attrs)
        cor = None
        m = re.search(r"color\s*:\s*([^;]+)", a.get("style") or "")
        if m:
            cor = m.group(1).strip()
        peso = re.search(r"font-weight\s*:\s*(bold|[6-9]00)", a.get("style") or "")
        self.pilha.append((tag, cor, bool(peso)))

    def handle_endtag(self, tag):
        if tag in self.BLOCO:
            self._fechar()
            return
        for i in range(len(self.pilha) - 1, -1, -1):
            if self.pilha[i][0] == tag:
                del self.pilha[i]
                break

    def _fmt(self):
        tags = {t for t, _, _ in self.pilha}
        cor = next((c for _, c, _ in reversed(self.pilha) if c), None)
        neg = bool(tags & {"b", "strong"}) or any(p for _, _, p in self.pilha)
        return (neg, bool(tags & {"i", "em"}), cor, "u" in tags)

    def handle_data(self, data):
        if not data:
            return
        if self.atual is None:
            if not data.strip():
                return
            self._novo("p", [])
        self.atual["runs"].append({"t": data, "fmt": self._fmt()})

    def _fechar(self):
        if self.atual is not None:
            self.blocos.append(self.atual)
        self.atual = None


def html_para_blocos(conteudo: str) -> list[dict]:
    lt = _Leitor()
    lt.feed(conteudo or "")
    lt._fechar()
    return lt.blocos


def texto_de_html(conteudo: str) -> str:
    return "\n".join("".join(r["t"] for r in b.get("runs") or []).strip()
                     for b in html_para_blocos(conteudo) if not b.get("quebra")).strip()


# ======================================================== exportacao
def _cor_rgb(cor: str):
    from docx.shared import RGBColor

    cor = (cor or "").strip().lower()
    m = re.fullmatch(r"#?([0-9a-f]{6})", cor)
    if m:
        return RGBColor.from_string(m.group(1).upper())
    m = re.fullmatch(r"rgb\((\d+),\s*(\d+),\s*(\d+)\)", cor)
    if m:
        return RGBColor(*(int(x) for x in m.groups()))
    return {"red": RGBColor(0xC0, 0, 0)}.get(cor)


def gerar_docx(conteudo: str, destino: Path, fonte: str = "Times New Roman", tamanho: float = 12,
               entrelinha: float = 1.5) -> Path:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH as AL, WD_BREAK
    from docx.shared import Cm, Pt

    d = docx.Document(str(TIMBRADO)) if TIMBRADO.exists() else docx.Document()  # sem timbrado: papel em branco
    corpo = d.element.body
    for el in list(corpo):
        if not el.tag.endswith("sectPr"):
            corpo.remove(el)
    for s in d.sections:  # margens que respeitam cabecalho e rodape do timbrado
        s.top_margin, s.bottom_margin = Cm(3.5), Cm(3.8)
        s.left_margin, s.right_margin = Cm(3.0), Cm(2.0)
    normal = d.styles["Normal"]
    normal.font.name, normal.font.size = fonte, Pt(tamanho)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts")
    if rfonts is not None:
        for k in list(rfonts.attrib):
            if k.endswith("Theme"):
                del rfonts.attrib[k]
    alin = {"center": AL.CENTER, "right": AL.RIGHT, "justify": AL.JUSTIFY, "left": AL.LEFT}
    quebra_pendente = False
    for b in html_para_blocos(conteudo):
        if b.get("quebra"):
            quebra_pendente = True
            continue
        p = d.add_paragraph()
        if quebra_pendente:
            p.paragraph_format.page_break_before = True
            quebra_pendente = False
        pf = p.paragraph_format
        pf.space_after = Pt(6)
        pf.line_spacing = entrelinha
        classe = b.get("classe")
        if classe == "citacao":
            pf.left_indent, pf.line_spacing = Cm(4), 1.0
        elif classe == "recuo":
            pf.first_line_indent = Cm(2)
        p.alignment = alin.get(b.get("alinha") or "", AL.JUSTIFY)
        for r in b.get("runs") or []:
            partes = r["t"].split("\n")
            for k, pedaco in enumerate(partes):
                if k:
                    p.add_run().add_break(WD_BREAK.LINE)
                if not pedaco:
                    continue
                run = p.add_run(pedaco)
                neg, ita, cor, sub = (list(r["fmt"]) + [False])[:4]
                run.bold = neg or classe == "titulo" or None
                run.italic = ita or None
                run.underline = sub or None
                run.font.name = fonte
                run.font.size = Pt(10 if classe == "citacao" else tamanho)
                rgb = _cor_rgb(cor) if cor else None
                if rgb is not None:
                    run.font.color.rgb = rgb
    d.save(str(destino))
    return destino


_PS_CONVERTE = r"""
param($in,$out)
$w = New-Object -ComObject Word.Application
$w.Visible = $false; $w.DisplayAlerts = 0
try { $d = $w.Documents.Open($in, $false, $true); $d.ExportAsFixedFormat($out, 17); $d.Close(0) }
finally { $w.Quit() }
"""


def docx_para_pdf(docx_path: Path, pdf_path: Path) -> Path:
    """Converte pelo Word (mesma paginacao do protocolo). Requer Microsoft Word instalado."""
    with tempfile.NamedTemporaryFile("w", suffix=".ps1", delete=False, encoding="utf-8") as f:
        f.write(_PS_CONVERTE)
        script = f.name
    try:
        r = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", script,
                            str(docx_path), str(pdf_path)], capture_output=True, text=True, timeout=180,
                           creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    finally:
        Path(script).unlink(missing_ok=True)
    if not pdf_path.exists():
        raise RuntimeError("Falha ao converter no Word: %s" % (r.stderr or r.stdout)[:300])
    return pdf_path


# ======================================================== IA: assistente de edicao
SYSTEM_EDITOR = (
    "Voce e um assistente de redacao juridica brasileira e edita uma peca processual a pedido do "
    "advogado responsavel. Recebe a peca em paragrafos numerados [§N], a instrucao e, quando houver, "
    "o trecho selecionado e o resultado da analise da peca (pontos fracos, criterios que faltam e "
    "textos sugeridos).\n"
    "Regras:\n"
    "- Faca SOMENTE o que foi pedido; preserve o restante, o estilo e a formatacao.\n"
    "- Linguagem forense formal, correta e objetiva, em portugues do Brasil.\n"
    "- Nao invente fatos, datas, valores, folhas, numeros de processo ou precedentes. Quando faltar um "
    "dado, escreva [COMPLETAR: o que falta]; para precedente a localizar, [LOCALIZAR PRECEDENTE].\n"
    "- Mantenha exatamente os marcadores ⟦D1⟧, ⟦D2⟧... (dados pessoais protegidos).\n"
    "- No texto pode usar <b>, <i> e <u>. Cada paragrafo novo e' uma operacao propria.\n"
    "- Se a instrucao for so uma pergunta, responda sem operacoes.\n"
    "Responda SOMENTE JSON:\n"
    '{"resposta":"o que foi feito ou a resposta, em ate 4 frases",'
    '"operacoes":[{"op":"substituir|inserir_depois|inserir_antes|remover","paragrafo":1,'
    '"texto":"texto completo do paragrafo","motivo":""}]}'
)

_DADO = re.compile(r"\b\d{3}\.?\d{3}\.?\d{3}-?\d{2}\b|\b\d{2}\.?\d{3}\.?\d{3}/?\d{4}-?\d{2}\b")


def _proteger(textos: list[str]) -> tuple[list[str], dict]:
    """Troca CPF/CNPJ por marcadores estaveis antes de ir ao provedor (e volta depois)."""
    mapa, inverso = {}, {}

    def sub(m):
        v = m.group(0)
        if v not in inverso:
            k = "⟦D%d⟧" % (len(inverso) + 1)
            inverso[v], mapa[k] = k, v
        return inverso[v]

    return [_DADO.sub(sub, t) for t in textos], mapa


def _restaurar(t: str, mapa: dict) -> str:
    for k, v in mapa.items():
        t = t.replace(k, v)
    return t


_PERMITIDAS = re.compile(r"</?(?!(?:b|i|u|strong|em)\b)[a-z][^>]*>", re.I)


def limpar_html_ia(t: str) -> str:
    t = _PERMITIDAS.sub("", str(t or ""))
    return re.sub(r"\s+\n|\n\s+", "\n", t).strip()


def contexto_analise(analise: dict | None, limite: int = 14000) -> str:
    if not analise:
        return "(a peca ainda nao foi analisada)"
    linhas = []
    for t in analise.get("tccs") or []:
        linhas.append({"topico": t.get("n"), "titulo": t.get("titulo"), "robustez": t.get("robustez"),
                       "tese_do_reu": (t.get("fundamento_réu") or "")[:300],
                       "faltam": [{"criterio": c.get("titulo"), "como_ganhar": c.get("como_ganhar")}
                                  for c in t.get("faltam") or []],
                       "texto_sugerido": (t.get("minuta_sugerida") or "")[:1200]})
    dec = analise.get("decisao") or {}
    extra = {"melhorias": [{"acao": m.get("acao"), "minuta": (m.get("minuta") or "")[:800]}
                           for m in dec.get("melhorias_prioritarias") or []][:6],
             "deficiencias": [x.get("ponto") for x in dec.get("deficiencias_replica") or []][:8]}
    return json.dumps({"topicos": linhas, **extra}, ensure_ascii=False)[:limite]


def editar(paragrafos: list[str], instrucao: str, selecao: dict | None, analise: dict | None,
           usar_analise: bool = True) -> dict:
    textos, mapa = _proteger(paragrafos)
    sel_txt = ""
    if selecao and selecao.get("texto"):
        sel_txt, _ = _proteger([str(selecao.get("texto"))[:4000]])
        sel_txt = "TRECHO SELECIONADO (paragrafos %s):\n%s\n\n" % (
            ", ".join(str(i + 1) for i in selecao.get("indices") or []), sel_txt[0])
    numerada = "\n\n".join("[§%d] %s" % (i + 1, t) for i, t in enumerate(textos))
    user = ("INSTRUCAO DO ADVOGADO: %s\n\n%s%s\n\n=== PECA ===\n%s" % (
        instrucao.strip()[:3000], sel_txt,
        ("RESULTADO DA ANALISE:\n" + contexto_analise(analise)) if usar_analise else "",
        numerada[:110000]))
    dados, modelo = ia.chat_json(SYSTEM_EDITOR, user, temperature=0.2, timeout=360)
    ops = []
    for o in dados.get("operacoes") or []:
        if not isinstance(o, dict):
            continue
        op = str(o.get("op") or "").strip().lower()
        try:
            n = int(o.get("paragrafo"))
        except (TypeError, ValueError):
            continue
        if op not in ("substituir", "inserir_depois", "inserir_antes", "remover") or not 1 <= n <= max(1, len(paragrafos)):
            continue
        texto = limpar_html_ia(_restaurar(str(o.get("texto") or ""), mapa))
        if op != "remover" and not texto:
            continue
        ops.append({"op": op, "paragrafo": n, "texto": texto, "motivo": str(o.get("motivo") or "")[:300]})
    return {"resposta": _restaurar(str(dados.get("resposta") or ""), mapa), "operacoes": ops, "modelo": modelo}


# ======================================================== avaliacao rapida (mesmo roteiro)
SYSTEM_AVALIA = (
    "Voce e um revisor tecnico de pecas juridicas brasileiras. Recebe os topicos do argumento da "
    "parte contraria e a peca de resposta em paragrafos [§N]. Para CADA topico, marque cada criterio "
    "do roteiro como sim, parcial ou nao, olhando so o texto da peca de resposta.\n\nROTEIRO:\n"
    + motor._roteiro + "\n\n'sim' so quando plenamente atendido; na duvida, 'parcial'. "
    "falta: o que ainda falta no criterio (curto). Responda SOMENTE JSON:\n"
    '{"topicos":[{"n":1,"criterios":{"enfrenta":{"nivel":"sim|parcial|nao","falta":""},"fatos":{},"lei":{},'
    '"juris":{},"prova":{},"consequencia":{},"coerencia":{}},"paragrafos":[1]}]}'
)


def avaliar_rapido(paragrafos: list[str], analise: dict) -> dict:
    tccs = [t for t in analise.get("tccs") or [] if t.get("status") != "erro"]
    textos, _ = _proteger(paragrafos)
    topicos = "\n".join("TOPICO %s: %s\nTese: %s" % (t["n"], t["titulo"], (t.get("fundamento_réu") or "")[:500])
                        for t in tccs)
    numerada = "\n\n".join("[§%d] %s" % (i + 1, re.sub(r"<[^>]+>", "", t)) for i, t in enumerate(textos))
    dados, modelo = ia.chat_json(SYSTEM_AVALIA, "TOPICOS:\n%s\n\n=== PECA DE RESPOSTA ===\n%s" % (
        topicos, numerada[:110000]), temperature=0.0, timeout=360)
    por_n = {}
    for x in dados.get("topicos") or []:
        if isinstance(x, dict):
            try:
                por_n[int(x.get("n"))] = x
            except (TypeError, ValueError):
                continue
    out, chances = [], []
    for t in tccs:
        x = por_n.get(int(t["n"])) or {}
        crit = x.get("criterios") if isinstance(x.get("criterios"), dict) else {}
        forca = {c["id"]: {"nivel": c["nivel"]} for c in t.get("criterios_reu") or []}
        if not crit:
            out.append({"n": t["n"], "robustez": t["robustez"], "chance_autor": t["nota_autor"], "sem_resposta": True})
            chances.append(t["nota_autor"])
            continue
        av = motor.avaliar_criterios({k: {"nivel": (v or {}).get("nivel") if isinstance(v, dict) else v,
                                          "como_chegar_a_sim": (v or {}).get("falta") if isinstance(v, dict) else ""}
                                      for k, v in crit.items()}, forca)
        forca_reu = t.get("forca_reu") if t.get("forca_reu") is not None else av["forca_reu"]
        ch = motor.chance_autor(av["robustez"], forca_reu)
        chances.append(ch)
        out.append({"n": t["n"], "robustez": av["robustez"], "chance_autor": ch,
                    "status": motor.status_de(av["criterios"], av["robustez"]),
                    "criterios": [{k: c[k] for k in ("id", "titulo", "peso", "nivel", "pontos", "como_ganhar")}
                                  for c in av["criterios"]],
                    "paragrafos": [p for p in x.get("paragrafos") or [] if isinstance(p, int)][:8]})
    calc = analise.get("calculo_probabilidade") or {}
    pesos = [t.get("peso_tese") or 1 for t in tccs]
    media = sum(c * w for c, w in zip(chances, pesos)) / sum(pesos) if chances else 50
    pa = int(max(5, min(95, round(media + float(calc.get("ajuste") or 0)))))
    return {"topicos": out, "probabilidade_autor": pa, "probabilidade_reu": 100 - pa, "modelo": modelo}
